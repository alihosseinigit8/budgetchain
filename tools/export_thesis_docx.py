"""Build the editable BSc thesis DOCX from the locally extracted thesis content.

The source data contains only the material authored for this project.  The
script deliberately produces a plain DOCX package instead of asking Word to
save through COM, because field updates in the supplied template can block an
unattended Word process.
"""

from __future__ import annotations

import json
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


WORKSPACE = Path(__file__).resolve().parents[1]
TEMPLATE = Path(os.environ.get(
    "THESIS_TEMPLATE",
    r"C:\Users\User\Downloads\BScProjectTemplate.docx",
))
CONTENT = WORKSPACE / "thesis_content_export.json"
OUTPUT = Path(os.environ.get(
    "THESIS_OUTPUT",
    str(WORKSPACE / "BudgetChain_BSc_Thesis_Draft.docx"),
))
UNIVERSITY_STYLE = os.environ.get("THESIS_STYLE", "standard").lower() == "university"
ASSET_DIR = WORKSPACE / "tools" / "thesis_template_assets"
BASMALLAH_IMAGE = ASSET_DIR / "image1.jpeg"
UNIVERSITY_LOGO = ASSET_DIR / "university-logo.png"


def set_run_font(run, name: str = "B Nazanin", size: float | None = None) -> None:
    run.font.name = name
    if size is not None:
        run.font.size = Pt(size)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn(f"w:{attr}"), name)
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), "fa-IR")
    lang.set(qn("w:bidi"), "fa-IR")
    complex_script = OxmlElement("w:cs")
    rpr.append(complex_script)
    rtl = OxmlElement("w:rtl")
    rpr.append(rtl)


def set_rtl(paragraph) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.append(bidi)


def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def border_cell(cell, color: str = "4F81BD", size: str = "14") -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "B Nazanin"
    normal.font.size = Pt(13)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.35

    for name, size, color in (
        ("Heading 1", 16, RGBColor(0, 0, 0)),
        ("Heading 2", 14, RGBColor(0, 0, 0)),
        ("Heading 3", 13, RGBColor(0, 0, 0)),
    ):
        style = document.styles[name]
        style.font.name = "B Lotus"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(14)
        style.paragraph_format.space_after = Pt(8)

    if "Caption" not in [style.name for style in document.styles]:
        caption = document.styles.add_style("Caption", WD_STYLE_TYPE.PARAGRAPH)
        caption.font.name = "B Nazanin"
        caption.font.size = Pt(11)


def clear_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def alignment_from_word(value: int) -> WD_ALIGN_PARAGRAPH:
    # Word values: 0 left, 1 centre, 2 right, 3 justify.
    return {
        0: WD_ALIGN_PARAGRAPH.LEFT,
        1: WD_ALIGN_PARAGRAPH.CENTER,
        2: WD_ALIGN_PARAGRAPH.RIGHT,
        3: WD_ALIGN_PARAGRAPH.JUSTIFY,
    }.get(value, WD_ALIGN_PARAGRAPH.RIGHT)


def normalize_persian(text: str) -> str:
    """Use a Word-safe Persian spelling for ezafe and normalize Arabic variants.

    Some Word/font combinations render the combining hamza in ``هٔ`` as a
    detached mark.  ``ه‌ی`` is a standard, readable digital spelling and
    avoids that rendering defect without changing the sentence meaning.
    """
    return (
        text.replace("\u0643", "ک")
        .replace("\u064a", "ی")
        .replace("هٔ", "ه‌ی")
    )


def add_centered_line(
    document: Document,
    text: str,
    *,
    font: str = "Zar",
    size: float = 14,
    bold: bool = False,
    before: float = 0,
    after: float = 0,
) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    set_rtl(paragraph)
    run = paragraph.add_run(normalize_persian(text))
    set_run_font(run, name=font, size=size)
    run.bold = bold


def add_university_front_matter(document: Document) -> None:
    """Recreate the Basmallah and cover-page visual language of the supplied template."""
    if not BASMALLAH_IMAGE.exists() or not UNIVERSITY_LOGO.exists():
        raise FileNotFoundError("The extracted Basmallah or University of Isfahan logo is missing.")

    # Page 1 — independent Basmallah page, as in the submitted visual template.
    basmallah = document.add_paragraph()
    basmallah.alignment = WD_ALIGN_PARAGRAPH.CENTER
    basmallah.paragraph_format.space_before = Cm(6.2)
    basmallah.add_run().add_picture(str(BASMALLAH_IMAGE), width=Cm(14.2))
    document.add_page_break()

    # Page 2 — cover.  The information comes from the thesis, while placement,
    # visual hierarchy, and font follow the supplied University of Isfahan sample.
    logo = document.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.paragraph_format.space_before = Pt(0)
    logo.paragraph_format.space_after = Pt(4)
    logo.add_run().add_picture(str(UNIVERSITY_LOGO), width=Cm(2.6))
    add_centered_line(document, "دانشگاه اصفهان", size=15, bold=True, after=3)
    add_centered_line(document, "دانشکدهٔ مهندسی کامپیوتر", size=14, bold=True, after=3)
    add_centered_line(document, "گروه مهندسی کامپیوتر", size=14, bold=True)
    add_centered_line(document, "پایان‌نامهٔ کارشناسی", size=19, bold=True, before=78, after=9)
    add_centered_line(document, "رشتهٔ مهندسی کامپیوتر ـ گرایش نرم‌افزار و شبکه", size=16, bold=True)
    add_centered_line(document, "عنوان پایان‌نامه", size=15, bold=True, before=64, after=9)
    add_centered_line(
        document,
        "طراحی و پیاده‌سازی نمونهٔ مفهومی شبکهٔ بلاک‌چینی مجوزدار برای مدیریت تخصیص اعتبار درون‌سازمانی",
        size=15,
        bold=True,
        after=62,
    )
    add_centered_line(document, "استاد راهنما: دکتر حمید [نام خانوادگی استاد راهنما]", size=14, after=5)
    add_centered_line(document, "پژوهشگر: علی حسینی‌فرد ـ شمارهٔ دانشجویی: ۴۰۰۳۶۲۳۰۱۳", size=13, after=5)
    add_centered_line(document, "تاریخ دفاع: [ماه و سال دفاع]", size=13)
    document.add_page_break()


def remove_generated_cover(events: list[dict]) -> None:
    """Replace the former plain text cover with the template-matched cover."""
    first_break = next(
        (index for index, event in enumerate(events)
         if event.get("type") == "paragraph" and event.get("text") == "\f"),
        None,
    )
    if first_break is None:
        raise ValueError("Could not identify the end of the generated cover page.")
    del events[:first_break + 1]


def add_toc_entry(document: Document, text: str) -> None:
    title, page = (normalize_persian(part.strip()) for part in text.split("\t", 1))
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(2.1), WD_TAB_ALIGNMENT.LEFT, WD_TAB_LEADER.DOTS
    )
    set_rtl(paragraph)
    title_run = paragraph.add_run(title)
    set_run_font(title_run, name="Zar", size=13)
    paragraph.add_run().add_tab()
    page_run = paragraph.add_run(page)
    set_run_font(page_run, name="Zar", size=13)


def add_list_heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    set_rtl(paragraph)
    run = paragraph.add_run(normalize_persian(text))
    set_run_font(run, name="Zar", size=16)
    run.bold = True


def add_paragraph(document: Document, event: dict) -> None:
    text = normalize_persian(event["text"])
    page_breaks = text.count("\f")
    text = text.replace("\f", "").strip("\r")
    for _ in range(page_breaks):
        document.add_page_break()
    if not text.strip():
        return

    outline = int(event.get("outline", 10))
    if outline == 1 and text.startswith("فصل "):
        chapter, separator, chapter_title = text.partition(":")
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.paragraph_format.space_before = Pt(22)
        paragraph.paragraph_format.space_after = Pt(12)
        paragraph.paragraph_format.page_break_before = True
        set_rtl(paragraph)
        chapter_run = paragraph.add_run(chapter)
        set_run_font(chapter_run, name="B Lotus", size=18)
        chapter_run.bold = True
        if separator and chapter_title.strip():
            chapter_run.add_break()
            title_run = paragraph.add_run(chapter_title.strip())
            set_run_font(title_run, name="B Lotus", size=18)
            title_run.bold = True
        return

    style = None
    if outline == 1:
        style = "Heading 1"
    elif outline == 2:
        style = "Heading 2"
    elif outline == 3:
        style = "Heading 3"
    elif text.startswith(("شکل ", "جدول ")):
        style = "Caption"

    paragraph = document.add_paragraph(style=style)
    paragraph.alignment = alignment_from_word(int(event.get("alignment", 2)))
    set_rtl(paragraph)
    if style == "Caption":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if outline == 1:
        paragraph.paragraph_format.page_break_before = True
    run = paragraph.add_run(text)
    heading_font = "B Lotus" if style in {"Heading 1", "Heading 2", "Heading 3"} else "B Nazanin"
    set_run_font(run, name=heading_font, size=11 if style == "Caption" else None)
    if style == "Caption":
        run.italic = True


def add_table(document: Document, event: dict) -> None:
    rows = event["rows"]
    if not rows:
        return
    cols = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.autofit = True
    single_cell = len(rows) == 1 and cols == 1
    for row_idx, row in enumerate(rows):
        for col_idx in range(cols):
            value = normalize_persian(row[col_idx]) if col_idx < len(row) else ""
            cell = table.cell(row_idx, col_idx)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if single_cell else WD_ALIGN_PARAGRAPH.RIGHT
            set_rtl(paragraph)
            run = paragraph.add_run(value)
            set_run_font(run, size=12 if single_cell else 11)
            if row_idx == 0 and not single_cell:
                run.bold = True
                shade(cell, "D9EAF7")
            if single_cell:
                border_cell(cell)
                shade(cell, "F5F9FC")
    document.add_paragraph()


def normalize_static_toc(events: list[dict]) -> None:
    """Keep the hand-authored contents list in the same order as the document.

    It is static by design: the original template's dynamic TOC field blocks
    unattended Word saves.  Page numbers come from the fully paginated source
    document and must be refreshed only if the student later changes content.
    """
    toc_entries = [
        "فصل اول: مقدمه\t12",
        "فصل دوم: مفاهیم و مبانی نظری\t16",
        "فصل سوم: تحلیل، طراحی و معماری پروژه\t22",
        "فصل چهارم: پیاده‌سازی، آزمون و ارزیابی\t31",
        "فصل پنجم: نتیجه‌گیری و پیشنهادهای آینده\t38",
        "پیوست الف: راهنمای اجرای نمونهٔ مفهومی\t41",
        "پیوست ب: ماتریس دادهٔ قابل انتشار\t43",
        "منابع\t45",
    ]
    start = next(
        (index for index, event in enumerate(events)
         if event.get("type") == "paragraph"
         and event.get("text", "").startswith("فصل اول: مقدمه\t")),
        None,
    )
    if start is not None:
        for index, text in enumerate(toc_entries):
            events[start + index]["text"] = text


def remove_english_abstract(events: list[dict]) -> None:
    """Keep this Persian thesis to one Persian abstract, as requested."""
    start = next(
        (index for index, event in enumerate(events)
         if event.get("type") == "paragraph" and event.get("text") == "Abstract"),
        None,
    )
    if start is None:
        return
    end = next(
        (index for index in range(start + 1, len(events))
         if events[index].get("type") == "paragraph"
         and events[index].get("text") == "فهرست مطالب"),
        None,
    )
    if end is None:
        raise ValueError("Could not identify the end of the English abstract section.")
    del events[start:end]


def main() -> None:
    if not TEMPLATE.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE}")
    if not CONTENT.exists():
        raise FileNotFoundError(f"Extracted content not found: {CONTENT}")
    if OUTPUT.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {OUTPUT}")

    payload = json.loads(CONTENT.read_text(encoding="utf-8-sig"))
    remove_english_abstract(payload["events"])
    normalize_static_toc(payload["events"])
    if UNIVERSITY_STYLE:
        remove_generated_cover(payload["events"])
    document = Document(TEMPLATE)
    clear_body(document)
    configure_styles(document)
    for section in document.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3)
        section.right_margin = Cm(2.5)

    if UNIVERSITY_STYLE:
        add_university_front_matter(document)

    toc_mode = False
    list_headings = {"فهرست شکل‌ها", "فهرست جدول‌ها", "فهرست اختصارات"}
    for event in payload["events"]:
        if event["type"] == "paragraph":
            event_text = event["text"]
            if UNIVERSITY_STYLE and event_text == "فهرست مطالب":
                add_list_heading(document, event_text)
                toc_mode = True
            elif UNIVERSITY_STYLE and toc_mode and "\t" in event_text and "\f" not in event_text:
                add_toc_entry(document, event_text)
            elif UNIVERSITY_STYLE and event_text in list_headings:
                add_list_heading(document, event_text)
                toc_mode = False
            else:
                add_paragraph(document, event)
                if event_text == "\f":
                    toc_mode = False
        elif event["type"] == "table":
            add_table(document, event)

    document.save(OUTPUT)
    print(f"THESIS_CREATED={OUTPUT}")
    print(f"EVENTS={len(payload['events'])}; SOURCE_PAGES={payload.get('pages', 'n/a')}")


if __name__ == "__main__":
    main()
